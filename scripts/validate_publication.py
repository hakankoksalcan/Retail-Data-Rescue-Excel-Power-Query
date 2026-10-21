"""Validate real source, reference exports, saved Excel results and publication links.
This program does not run Excel or evaluate Power Query M.
Run the reproduction commands in README.md first, then this file.
"""
from pathlib import Path
from datetime import datetime,timezone
from decimal import Decimal
from zipfile import ZipFile
from numbers import Number
from urllib.parse import unquote,urlsplit
import base64,csv,gzip,hashlib,io,json,re,struct
from collections import Counter
import openpyxl,pandas as pd,fitz
from lxml import etree as ET

ROOT=Path(__file__).resolve().parents[1]
results=[]
def check(name,expected,actual,scope='Publication artifacts'):
 good=abs(expected-actual)<=0.0001 if isinstance(expected,float) and isinstance(actual,Number) else expected==actual
 results.append({'Control':name,'Expected':expected,'Actual':actual,'Status':'PASS' if good else 'FAIL','Scope':scope})
def digest(path):return hashlib.sha256(path.read_bytes()).hexdigest()

raw=ROOT/'raw'/'Online Retail.xlsx'
check('Original source SHA256','43465a06f2ccf7c8b5bd2892bc7defb52f97487934fe93b16ae4c3936424676d',digest(raw),'Original source')
source=pd.read_excel(raw,sheet_name='Online Retail',dtype=object)
fields=['InvoiceNo','StockCode','Description','Quantity','InvoiceDate','UnitPrice','CustomerID','Country']
check('Original column schema',fields,list(source.columns),'Original source')
check('Original row count',541909,len(source),'Original source')
missing={k:int(source[k].isna().sum()) for k in fields}
check('Original null counts',{'InvoiceNo':0,'StockCode':0,'Description':1454,'Quantity':0,'InvoiceDate':0,'UnitPrice':0,'CustomerID':135080,'Country':0},missing,'Original source')
check('Native source dates',541909,int(source.InvoiceDate.map(lambda v:isinstance(v,datetime)).sum()),'Original source')
check('Whole-number native quantities',541909,int(source.Quantity.map(lambda v:isinstance(v,Number) and float(v).is_integer()).sum()),'Original source')
check('Native numeric source prices',541909,int(source.UnitPrice.map(lambda v:isinstance(v,Number)).sum()),'Original source')
check('Observed first timestamp','2010-12-01T08:26:00',min(source.InvoiceDate).isoformat(),'Original source')
check('Observed last timestamp','2011-12-09T12:50:00',max(source.InvoiceDate).isoformat(),'Original source')
check('Additional exact records',5268,int(source.duplicated(fields,keep='first').sum()),'Original source')
check('Zero price rows',2515,int(source.UnitPrice.eq(0).sum()),'Original source')
check('Negative price rows',2,int(source.UnitPrice.lt(0).sum()),'Original source')
raw_units=sum(int(Decimal(str(p))*10000)*int(q) for p,q in zip(source.UnitPrice,source.Quantity))
check('Independent raw integer amount',97477479340,raw_units,'Original source')

summary=json.loads((ROOT/'reference'/'summary.json').read_text());totals=summary['totals']
ref_controls=pd.read_csv(ROOT/'reference'/'Controls.csv')
check('Executed reference failures',0,int(ref_controls.Status.eq('FAIL').sum()),'Python reference')
check('Executed reference passes',30,int(ref_controls.Status.eq('PASS').sum()),'Python reference')
check('Explicit full-model Excel acceptance pending',1,int(ref_controls.Status.eq('PENDING').sum()),'Status coverage')
reconciliation=pd.read_csv(ROOT/'reference'/'Reconciliation.csv')
check('Class rows reconcile',len(source),int(reconciliation.Rows.sum()),'Python reference')
check('Class amounts reconcile in integer units',raw_units,int(reconciliation.AmountUnits4.sum()),'Python reference')
expected_classes={'MerchandiseSale':527761,'MerchandiseCancellation':8704,'SpecialOperation':2946,'ZeroPriceAdjustment':1335,'UnresolvedReview':1163}
check('Five exclusive class row counts',expected_classes,dict(zip(reconciliation.TransactionClass,reconciliation.Rows)),'Python reference')

book=openpyxl.load_workbook(ROOT/'Retail_Data_Rescue.xlsx',data_only=True)
formulas=openpyxl.load_workbook(ROOT/'Retail_Data_Rescue.xlsx',data_only=False)
check('Workbook opens on Dashboard','Dashboard',book.active.title)
check('First worksheet Dashboard','Dashboard',book.sheetnames[0])
check('Settings table retained',['tblSource'],list(book['Settings'].tables))
check('Source path configured sample',r'C:\RetailDataRescue\raw\Online Retail.xlsx',book['Settings']['A6'].value)
check('SourceCheck saved results',[("Kontrol","Sonuç"),("Veri satırı sayısı",541909),("Sütun sayısı",8)],list(book['SourceCheck'].values),'Saved desktop Excel audit')
check('Missing-values saved results',missing,{row[0]:row[1] for row in list(book['qMissingValues'].values)[1:]},'Saved desktop Excel audit')
check('Dashboard default filters',['All','2010-12','2011-12'],[book['Dashboard'][k].value for k in ['B5','F5','J5']])
for address,metric in [('B9','ReportableSalesGBP'),('F9','MerchandiseCancellationGBP'),('J9','MerchandiseBalanceGBP'),('B13','EligibleSalesInvoiceIDs'),('F13','KnownCustomerRowCoverage'),('J13','ReviewRows')]:
 check('Saved dashboard '+metric,totals[metric],book['Dashboard'][address].value,'Saved formula result vs fresh reference')
 check('Live Excel formula '+address,True,formulas['Dashboard'][address].data_type=='f')
check('Scoped status visible',True,'full-model integration pending' in book['RefreshInfo']['B2'].value)
errors=[s.title+'!'+c.coordinate+':'+str(c.value) for s in book if s.title!='ChartData' for row in s for c in row if c.data_type=='e']
check('Unexpected visible formula errors',[],errors)
check('Dashboard chart named ranges',['rdrCancels','rdrMonths','rdrSales'],sorted(book['Dashboard'].defined_names))
check('Native dashboard chart count',3,len(formulas['Dashboard']._charts))
# Every numeric reporting table is compared with a freshly generated reference table.
for name in ['KPI','Monthly','Country','Product','Reconciliation','Bridge','Quality','Controls','DuplicateScenario','ProductDictionary','SpecialOperations','Customers','DashboardData']:
 data=pd.read_csv(ROOT/'reference'/(name+'.csv'),keep_default_na=False,dtype={'CustomerID':str,'StockCode':str})
 rows=list(book[name].values)
 headers=list(rows[0])
 check('Table row count '+name,len(data),len(rows)-1)
 for col in ['CustomerID','StockCode']:
  if col in data.columns and col in headers:
   saved_ids=[row[headers.index(col)] for row in rows[1:]]
   check('Identifier values '+name+'.'+col,True,saved_ids==data[col].tolist())
 numeric_columns=[col for col in data.columns if pd.api.types.is_numeric_dtype(data[col]) and col in headers]
 for col in numeric_columns:
  ix=headers.index(col)
  saved=[row[ix] for row in rows[1:]]
  generated=data[col].tolist()
  # Tables have a fixed order in the published reference generator.
  good=len(saved)==len(generated) and all(isinstance(a,Number) and abs(float(a)-float(b))<=0.0001 for a,b in zip(saved,generated))
  check('Table values '+name+'.'+col,True,good)

with ZipFile(ROOT/'Retail_Data_Rescue.xlsx') as z:
 check('XLSX archive CRC',None,z.testzip())
 connections=ET.fromstring(z.read('xl/connections.xml'))
 locations=sorted(re.search(r'Location=([^;]+)',node.get('connection')).group(1) for node in connections.xpath('//*[local-name()="dbPr"]'))
 expected_queries=['qMissingValues','qRawSource','qSourceCheck','qTypedSource']
 check('Embedded query connection names',expected_queries,locations)
 mash=base64.b64decode(ET.fromstring(z.read('customXml/item1.xml')).text)
 size=struct.unpack_from('<I',mash,4)[0]
 with ZipFile(io.BytesIO(mash[8:8+size])) as package:
  section=package.read('Formulas/Section1.m').decode('utf-8-sig').replace('\r\n','\n')
 exports={m.group(1):m.group(2).strip().removesuffix(';')+'\n' for m in re.finditer(r'^shared\s+(\w+)\s*=\s*(.*?)(?=^shared\s|\Z)',section,re.M|re.S)}
 check('Embedded formula names',expected_queries,sorted(exports))
 for name,code in exports.items():check('Exact M export '+name,code,(ROOT/'power_query'/'embedded'/(name+'.pq')).read_text())
 check('Embedded query table count',2,len([n for n in z.namelist() if n.startswith('xl/queryTables/queryTable') and n.endswith('.xml')]))
 check('Local Desktop presentation path removed',False,b'Users\\hakan\\Desktop' in z.read('xl/workbook.xml'))

refresh=json.loads((ROOT/'validation'/'refresh-tests.json').read_text())
check('Reference refresh check count',19,len(refresh['checks']),'Python reference refresh')
check('Reference refresh failures',0,sum(x['Status']=='FAIL' for x in refresh['checks']),'Python reference refresh')
syntax=json.loads((ROOT/'validation'/'m-syntax.json').read_text())
check('M files parsed',24,len(syntax['results']),'M syntax only')
check('M syntax failures',0,sum(x['status']=='FAIL' for x in syntax['results']),'M syntax only')
check('M parser evidence matches current files',True,all(x.get('sha256')==digest(ROOT/x['file']) for x in syntax['results']),'M syntax only')
check('Preserved query package',True,json.loads((ROOT/'validation'/'workbook-preservation.json').read_text())['power_query_package_preserved'])

ids=set();review_count=0;outside=0;badflag=0
with gzip.open(ROOT/'review'/'ReviewRecords.csv.gz','rt',encoding='utf-8-sig',newline='') as f:
 for row in csv.DictReader(f):
  review_count+=1;ids.add(row['RecordID'])
  n=int(row['SourceRow'])
  outside+=not(2<=n<=len(source)+1 and row['RecordID']==f'Online Retail.xlsx|Online Retail|{n:09d}')
  badflag+=row['NeedsReview']!='True'
check('Review register rows',232181,review_count,'Review export')
check('Review register unique RecordID',232181,len(ids),'Review export')
check('Review register source membership',0,outside,'Review export')
check('Review register flagged records',0,badflag,'Review export')
for name in ['Client_Handover.pdf','Dashboard_Reference.pdf']:
 with fitz.open(ROOT/name) as doc:check('One-page '+name,1,len(doc))

bad_links=[]
for p in [ROOT/'README.md',*(ROOT/'docs').glob('*.md'),ROOT/'power_query'/'README_TR.md']:
 for href in re.findall(r'!?\[[^\]]*\]\(([^)]+)\)',p.read_text()):
  target=href.strip().split(' "')[0]
  if urlsplit(target).scheme or target.startswith('#'):continue
  path=unquote(target.split('#')[0])
  if path and not (p.parent/path).exists():bad_links.append(str(p.relative_to(ROOT))+': '+target)
check('Broken relative documentation links',[],bad_links)
manifest=json.loads((ROOT/'FILE_MANIFEST.json').read_text())
bad_hashes=[row['path'] for row in manifest['files'] if not (ROOT/row['path']).is_file() or digest(ROOT/row['path'])!=row['sha256']]
check('Manifest hash mismatches',[],bad_hashes)
check('No file exceeds GitHub regular limit',True,all(row['bytes']<100*1024*1024 for row in manifest['files']))
check('No dependencies/caches/test XLSX in manifest',[],[row['path'] for row in manifest['files'] if any(x in row['path'].split('/') for x in ['.cache','node_modules','__pycache__']) or 'TEST_ONLY_' in row['path']])
check('No initial transfer helpers in manifest',[],[row['path'] for row in manifest['files'] if row['path'].startswith('delivery_downloads/')])

out={'checked_at_utc':datetime.now(timezone.utc).isoformat(),'scope':'Source/artifact/reference validation; no full-model Excel runtime execution','checks':results,'all_executed_checks_passed':all(x['Status']=='PASS' for x in results),'full_reporting_model_excel_acceptance':'PENDING','desktop_excel_audit_evidence':'User-supplied workbook and confirmations'}
(ROOT/'validation'/'publication-validation.json').write_text(json.dumps(out,indent=2,default=str)+'\n')
with (ROOT/'validation'/'publication-controls.csv').open('w',newline='',encoding='utf-8-sig') as f:
 writer=csv.DictWriter(f,lineterminator='\n',fieldnames=['Control','Expected','Actual','Status','Scope']);writer.writeheader();writer.writerows(results)
print('Publication checks:',len(results),'PASS:',sum(x['Status']=='PASS' for x in results))
for row in results:
 if row['Status']=='FAIL':print(json.dumps(row,default=str))
if not out['all_executed_checks_passed']:raise SystemExit(1)
