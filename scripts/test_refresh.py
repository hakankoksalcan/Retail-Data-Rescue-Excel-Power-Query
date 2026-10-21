"""Re-read same source, then a clearly labelled full-size test copy with one added row."""
from pathlib import Path
import hashlib,json,subprocess,sys
from decimal import Decimal
import openpyxl,pandas as pd

ROOT=Path(__file__).resolve().parents[1]
raw=ROOT/'raw'/'Online Retail.xlsx'
before=hashlib.sha256(raw.read_bytes()).hexdigest()
base=json.loads((ROOT/'reference'/'summary.json').read_text())
repeat=ROOT/'tests'/'repeat_reference'
subprocess.run([sys.executable,str(ROOT/'scripts'/'build_reference.py'),'--output',str(repeat),'--no-fact-export'],check=True)
repeat_summary=json.loads((repeat/'summary.json').read_text())
checks=[]
def record(name,expected,actual):
 good=expected==actual
 checks.append({'Control':name,'Expected':expected,'Actual':actual,'Status':'PASS' if good else 'FAIL','Scope':'Python reference refresh'})
record('Same-source repeated totals',base['totals'],repeat_summary['totals'])
record('Same-source repeated class partitions',base['classes'],repeat_summary['classes'])
for name in ['KPI','Reconciliation','Monthly','Country','Product','ProductDictionary','Quality','DuplicateScenario','Bridge']:
 record('Repeat table '+name,True,(ROOT/'reference'/(name+'.csv')).read_bytes()==(repeat/(name+'.csv')).read_bytes())

testpath=ROOT/'tests'/'TEST_ONLY_Online_Retail_plus_one.xlsx'
print('Writing separate TEST_ONLY workbook; raw stays untouched',flush=True)
src=openpyxl.load_workbook(raw,read_only=True,data_only=False)
dst=openpyxl.Workbook(write_only=True); sheet=dst.create_sheet('Online Retail')
source=src['Online Retail']; invoice_present=False
for row in source.iter_rows(values_only=True):
 if row[0] in [990001,'990001']:invoice_present=True
 sheet.append(row)
src.close()
assert not invoice_present,'Synthetic test invoice already exists in raw source'
from datetime import datetime
added=[990001,'85123A','WHITE HANGING HEART T-LIGHT HOLDER',2,datetime(2011,12,9,13,0),3.5,17850,'United Kingdom']
sheet.append(added); dst.save(testpath)
testout=ROOT/'tests'/'appended_reference'
subprocess.run([sys.executable,str(ROOT/'scripts'/'build_reference.py'),'--input',str(testpath),'--output',str(testout),'--no-fact-export'],check=True)
test=json.loads((testout/'summary.json').read_text())
record('Appended source row count',base['rows']+1,test['rows'])
record('Raw amount increases GBP 7.0000',Decimal('7'),Decimal(str(test['totals']['RawSignedAmountGBP']))-Decimal(str(base['totals']['RawSignedAmountGBP'])))
record('Sales amount increases GBP 7.0000',Decimal('7'),Decimal(str(test['totals']['ReportableSalesGBP']))-Decimal(str(base['totals']['ReportableSalesGBP'])))
record('Cancellation unchanged',base['totals']['MerchandiseCancellationGBP'],test['totals']['MerchandiseCancellationGBP'])
record('Sales invoice count increases one',base['totals']['EligibleSalesInvoiceIDs']+1,test['totals']['EligibleSalesInvoiceIDs'])
record('Known-customer sales rows increases one',base['totals']['KnownCustomerSalesRows']+1,test['totals']['KnownCustomerSalesRows'])
record('Duplicate extras unchanged',base['totals']['DuplicateExtraRows'],test['totals']['DuplicateExtraRows'])
record('Raw source SHA256 unchanged',before,hashlib.sha256(raw.read_bytes()).hexdigest())
manifest={'test_only':True,'synthetic_row':[{str(k):str(v) for k,v in zip(['InvoiceNo','StockCode','Description','Quantity','InvoiceDate','UnitPrice','CustomerID','Country'],added)}],
          'added_source_row':541911,'test_file':testpath.name,'checks':checks,
          'excel_refresh_status':'PENDING: repeat these steps in Excel',
          'all_python_checks_passed':all(c['Status']=='PASS' for c in checks)}
(ROOT/'validation'/'refresh-tests.json').write_text(json.dumps(manifest,indent=2,default=str),encoding='utf-8')
pd.DataFrame([{k:str(v) if isinstance(v,(dict,Decimal)) else v for k,v in c.items()} for c in checks]).to_csv(ROOT/'validation'/'refresh-controls.csv',index=False,encoding='utf-8-sig')
print(json.dumps(manifest,indent=2,default=str),flush=True)
if not manifest['all_python_checks_passed']:raise SystemExit('Refresh reference test failed')
